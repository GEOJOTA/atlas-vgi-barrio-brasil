#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
generar_doc_estado_avance.py
============================
Informe de Estado de Avance (corte 150 horas) de la Práctica Profesional
"Atlas de Información Geográfica Voluntaria para la Gestión Territorial".

Estructura: cada Objetivo Específico se desarrolla como una cadena
    OBJETIVO  →  PROCEDIMIENTOS (plan metodológico)  →  PRODUCTOS
usando exclusivamente los objetivos declarados en la presentación de avance
y las actividades / productos declarados en el Plan de Trabajo formal
(Planificación geopráctica profesional 2026.xlsx).

Salida:
/Users/geojplescano/Programas/Práctica profesional 2026 S2/Estado_Avance_Atlas_VGI_UAH.docx
"""

import os
import textwrap

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch, Circle

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import nsdecls
from docx.oxml import parse_xml

BASE_DIR = "/Users/geojplescano/Programas/Práctica profesional 2026 S2"
OUTPUT_DOCX = os.path.join(BASE_DIR, "Estado_Avance_Atlas_VGI_UAH.docx")
DIR_DIAGRAMAS = os.path.join(BASE_DIR, "scratch", "diagramas_informe")
os.makedirs(DIR_DIAGRAMAS, exist_ok=True)

# ---------------------------------------------------------------------------
# Paleta institucional
# ---------------------------------------------------------------------------
HEX_AZUL_UAH = "#003366"
HEX_AZUL_SEC = "#0F4C81"
HEX_AZUL_CLARO = "#E8F0F8"
HEX_GRIS_LINEA = "#94A3B8"
HEX_NEGRO = "#111827"

ESTADOS = {
    "completado": ("Completado", "#10B981"),
    "en_curso": ("En ejecución", "#F59E0B"),
    "anticipado": ("Avance anticipado", "#38BDF8"),
    "programado": ("Programado", "#94A3B8"),
}

COLOR_AZUL_UAH = RGBColor(0, 51, 102)
COLOR_AZUL_SEC = RGBColor(15, 76, 129)
COLOR_NEGRO = RGBColor(17, 24, 39)
COLOR_GRIS = RGBColor(100, 116, 139)
COLOR_BLANCO = RGBColor(255, 255, 255)

# ---------------------------------------------------------------------------
# CONTENIDO — Objetivo general y objetivos específicos (texto declarado en PPT)
# ---------------------------------------------------------------------------
OBJETIVO_GENERAL = (
    "Desarrollar un prototipo de Atlas de Información Geográfica Voluntaria (VGI) centrado en los usuarios "
    "para Barrio Brasil y el entorno del campus de la Universidad Alberto Hurtado, articulando datos colaborativos "
    "y fuentes oficiales, con el propósito de subsanar vacíos de información territorial a escala barrial y proveer "
    "herramientas interactivas de diagnóstico, análisis espacial y gestión urbana local."
)

OBJETIVOS = [
    {
        "id": "OE 1",
        "titulo": "Caracterización territorial y arquitectura conceptual",
        "corto": "Caracterización y\narquitectura conceptual",
        "texto": (
            "Caracterizar la disponibilidad de información geográfica voluntaria y oficial en Barrio Brasil, mediante "
            "la revisión sistemática de fuentes abiertas e institucionales, para estructurar la arquitectura conceptual "
            "del Atlas a partir de una taxonomía temática normalizada y perfiles territoriales centrados en el usuario."
        ),
        "hitos": "Hitos 1, 2 y 3",
        "hitos_corto": "H1–H3",
        "periodo": "Días 1–12 · 01/09 – 16/09/2026",
        "semanas": "Sem 1–3",
        "horas": 72,
        "estado": "completado",
        "procedimientos": [
            {
                "id": "P1.1", "hito": "H1",
                "nombre": "Diagnóstico institucional y delimitación del caso",
                "corto": "Diagnóstico institucional y delimitación",
                "detalle": (
                    "Se sostuvo la reunión de inducción con el referente del centro de práctica para conocer qué "
                    "información territorial utiliza y cuál le hace falta para trabajar el barrio. A partir de esa "
                    "conversación se delimitó el área de estudio (Barrio Brasil y entorno del campus UAH), se "
                    "identificaron los usuarios potenciales y se formularon la pregunta central, los objetivos y el alcance."
                ),
                "tecnica": "Entrevista / consulta institucional · revisión documental",
                "productos": ["1.1"],
            },
            {
                "id": "P1.2", "hito": "H1",
                "nombre": "Revisión bibliográfica sobre VGI",
                "corto": "Revisión bibliográfica VGI",
                "detalle": (
                    "Se leyeron y sistematizaron los textos de referencia sobre información geográfica voluntaria "
                    "(Goodchild, 2007; Haklay, 2010; Neis & Zipf, 2012) para entender cómo se produce el dato "
                    "colaborativo, cuáles son sus sesgos de cobertura y con qué criterios se evalúa su calidad."
                ),
                "tecnica": "Análisis documental y fichaje bibliográfico",
                "productos": ["1.1", "1.3"],
            },
            {
                "id": "P1.3", "hito": "H2",
                "nombre": "Definición de la taxonomía temática y de los perfiles de usuario",
                "corto": "Taxonomía y perfiles de usuario",
                "detalle": (
                    "Se revisaron las etiquetas de OpenStreetMap (OSM Wiki) y el estándar Darwin Core de GBIF para "
                    "agrupar los elementos urbanos en categorías temáticas útiles para la gestión del barrio. Con ello "
                    "se fijaron las 10 macrocategorías del Atlas (CAT-01 a CAT-10) y se describieron los perfiles de "
                    "usuario que orientan qué información debe mostrar el prototipo."
                ),
                "tecnica": "Análisis de taxonomías OSM / GBIF · diseño centrado en el usuario",
                "productos": ["1.3"],
            },
            {
                "id": "P1.4", "hito": "H2",
                "nombre": "Diseño de fichas de capa y plantillas de metadatos",
                "corto": "Fichas de capa y metadatos",
                "detalle": (
                    "Para cada categoría se elaboró una ficha con atributos mínimos, escala de representación, "
                    "indicadores temáticos y criterios de calidad, junto con una plantilla de metadatos ISO 19115. "
                    "Sobre esa base se bosquejó la estructura del visor y los criterios de interoperabilidad (GeoJSON, OGC/WMS)."
                ),
                "tecnica": "Modelamiento conceptual de datos · ISO 19115",
                "productos": ["1.3"],
            },
            {
                "id": "P1.5", "hito": "H3",
                "nombre": "Inventario y evaluación de fuentes VGI y oficiales",
                "corto": "Inventario de fuentes VGI",
                "detalle": (
                    "Se inventariaron las fuentes disponibles para el barrio (OSM, GBIF, ohsome, IDE Chile, MINVU "
                    "GeoIDE) y se evaluó cada una según tipo, cobertura, actualización, escala, licencia y utilidad. "
                    "El contraste entre lo que ofrece el dato oficial y lo que aporta el dato voluntario permitió "
                    "identificar las brechas de información que el Atlas debe cubrir."
                ),
                "tecnica": "Matriz multicriterio de fuentes · análisis de brechas",
                "productos": ["1.2"],
            },
        ],
        "productos": [
            {"id": "1.1", "nombre": "Ficha del proyecto", "corto": "Ficha del proyecto",
             "contenido": "Territorio, usuarios, objetivos, alcance y productos comprometidos.",
             "estado": "completado"},
            {"id": "1.2", "nombre": "Catálogo de fuentes VGI", "corto": "Catálogo de fuentes VGI",
             "contenido": "Inventario evaluado por tipo, cobertura, actualización, escala, licencia y utilidad.",
             "estado": "completado"},
            {"id": "1.3", "nombre": "Diseño conceptual del Atlas", "corto": "Diseño conceptual Atlas",
             "contenido": "Arquitectura de capas, 10 categorías OSM/GBIF, perfiles de usuario, atributos mínimos, "
                          "indicadores y plantillas de metadatos ISO 19115.",
             "estado": "completado"},
        ],
    },
    {
        "id": "OE 2",
        "titulo": "Geodatabase espacial, calidad y análisis territorial",
        "corto": "Geodatabase, calidad\ny análisis territorial",
        "texto": (
            "Estructurar una geodatabase espacial multicapa y desarrollar el análisis territorial de Barrio Brasil, "
            "aplicando depuración topológica, auditoría de calidad de datos bajo estándares internacionales y análisis "
            "espacial multicriterio para evaluar la relación entre los equipamientos urbanos y la población del área de estudio."
        ),
        "hitos": "Hitos 4, 5 y 6",
        "hitos_corto": "H4–H6",
        "periodo": "Días 13–32 · 21/09 – 19/10/2026",
        "semanas": "Sem 4–7",
        "horas": 120,
        "estado": "en_curso",
        "procedimientos": [
            {
                "id": "P2.1", "hito": "H4",
                "nombre": "Descarga automatizada de las fuentes seleccionadas",
                "corto": "Descarga automatizada de datos",
                "detalle": (
                    "Se programaron scripts en Python que consultan directamente Overpass API (OSM), ohsome API "
                    "(HeiGIT), GBIF API y los servicios de IDE Chile (MINVU GeoIDE, SUBDERE, ArcGIS Online), además de "
                    "los datos del Censo INE 2024. Así la descarga queda documentada y puede repetirse en cualquier momento."
                ),
                "tecnica": "Python (requests) · Overpass QL · APIs REST",
                "productos": ["2.1"],
            },
            {
                "id": "P2.2", "hito": "H4",
                "nombre": "Limpieza, depuración topológica y clasificación temática",
                "corto": "Limpieza y corrección topológica",
                "detalle": (
                    "Los datos descargados se reproyectaron a UTM 19S (EPSG:32719) para los cálculos métricos, se "
                    "eliminaron duplicados, se normalizaron los nombres de atributos y se repararon geometrías inválidas "
                    "(auto-intersecciones, anillos abiertos). Luego cada objeto se asignó a su categoría temática."
                ),
                "tecnica": "GeoPandas · Shapely · QGIS · PostGIS (ST_MakeValid)",
                "productos": ["2.1"],
            },
            {
                "id": "P2.3", "hito": "H4",
                "nombre": "Integración en la geodatabase y documentación",
                "corto": "Geodatabase e integración",
                "detalle": (
                    "Las capas depuradas se integraron en una geodatabase única (GeoJSON en EPSG:4326 y PostgreSQL/PostGIS), "
                    "organizada por categoría temática. Cada capa se documentó con su ficha de metadatos ISO 19115 y se "
                    "consolidaron los catálogos en Excel."
                ),
                "tecnica": "PostgreSQL / PostGIS · ISO 19115 · openpyxl",
                "productos": ["2.1", "2.4"],
            },
            {
                "id": "P2.4", "hito": "H5",
                "nombre": "Evaluación de calidad y cobertura del dato VGI",
                "corto": "Evaluación de calidad y cobertura",
                "detalle": (
                    "Se midió la completitud de atributos por categoría, la evolución temporal de los objetos (intervalos "
                    "anuales), la densidad por km², la intensidad de contribución (usuarios y ediciones), la saturación del "
                    "mapeo y la consistencia geométrica, revisando además los changesets con OSMCha para detectar errores o vandalismo."
                ),
                "tecnica": "ohsome API (/elements, /contributions, /users) · OSMCha · ISO 19157",
                "productos": ["2.3"],
            },
            {
                "id": "P2.5", "hito": "H5",
                "nombre": "Construcción del Índice de Calidad VGI (ICV)",
                "corto": "Índice de Calidad VGI (ICV)",
                "detalle": (
                    "Los indicadores anteriores se normalizaron y combinaron mediante una suma lineal ponderada "
                    "(completitud, actualidad, consistencia geométrica y densidad), obteniendo un puntaje comparable por "
                    "categoría temática y por capa, que indica qué tan confiable es cada conjunto de datos."
                ),
                "tecnica": "Combinación lineal ponderada (WLC) · normalización min–max",
                "productos": ["2.3"],
            },
            {
                "id": "P2.6", "hito": "H6",
                "nombre": "Análisis territorial equipamiento–población",
                "corto": "Análisis espacial y cruce INE",
                "detalle": (
                    "Se elaboran mapas de densidad kernel y de cobertura por categoría, un análisis centro–borde del barrio "
                    "y la relación entre objetos VGI y población (objetos por habitante, INE 2024), para identificar vacíos "
                    "territoriales y ordenar los sectores mediante estadísticas espaciales y rankings."
                ),
                "tecnica": "Densidad kernel · estadística espacial · cruce con Censo INE 2024",
                "productos": ["2.2"],
            },
        ],
        "productos": [
            {"id": "2.1", "nombre": "Geodatabase VGI de Barrio Brasil", "corto": "Geodatabase Barrio Brasil",
             "contenido": "Capas en GeoJSON (EPSG:4326) y PostGIS, organizadas por categoría temática.",
             "estado": "completado"},
            {"id": "2.2", "nombre": "Mapas analíticos", "corto": "Mapas analíticos",
             "contenido": "Densidad, cobertura por categoría, distribución centro–borde y relación VGI–población.",
             "estado": "en_curso"},
            {"id": "2.3", "nombre": "Índice de Calidad VGI (ICV)", "corto": "Índice de Calidad (ICV)",
             "contenido": "Puntaje compuesto por categoría temática y por capa.",
             "estado": "completado"},
            {"id": "2.4", "nombre": "Catálogos Excel", "corto": "Catálogos en Excel",
             "contenido": "Fuentes VGI, categorías OSM/GBIF y perfil de metadatos ISO 19115.",
             "estado": "completado"},
        ],
    },
    {
        "id": "OE 3",
        "titulo": "Prototipo cartográfico y geoportal web interactivo",
        "corto": "Prototipo cartográfico\ny geoportal web",
        "texto": (
            "Diseñar e implementar el prototipo cartográfico y la plataforma web interactiva del Atlas mediante una "
            "arquitectura abierta e interoperable, con el propósito de facilitar la visualización dinámica, la consulta "
            "de condiciones normativas, la conexión a servicios espaciales y la transferencia de datos a la comunidad."
        ),
        "hitos": "Hitos 7 y 8",
        "hitos_corto": "H7–H8",
        "periodo": "Días 33–44 · 20/10 – 04/11/2026",
        "semanas": "Sem 8–10",
        "horas": 72,
        "estado": "anticipado",
        "procedimientos": [
            {
                "id": "P3.1", "hito": "H7",
                "nombre": "Revisión de referentes cartográficos y de usabilidad",
                "corto": "Revisión referentes y usabilidad",
                "detalle": (
                    "Se revisan las recomendaciones de diseño cartográfico de la Asociación Cartográfica Internacional (ICA), "
                    "geoportales de referencia (IDE Chile, OSM, GBIF) y criterios de usabilidad web (Nielsen Norman Group), "
                    "para decidir cómo deben verse y leerse los mapas del Atlas."
                ),
                "tecnica": "Análisis comparativo de geoportales · heurísticas de usabilidad",
                "productos": ["3.1"],
            },
            {
                "id": "P3.2", "hito": "H7",
                "nombre": "Definición de la identidad gráfica del Atlas",
                "corto": "Identidad gráfica y simbología",
                "detalle": (
                    "Se fija una paleta de colores por categoría temática, la simbología de puntos, líneas y polígonos, "
                    "las tipografías y la composición de cada mapa, de modo que todas las capas compartan un mismo lenguaje visual."
                ),
                "tecnica": "Diseño cartográfico · semiología gráfica",
                "productos": ["3.1", "3.2"],
            },
            {
                "id": "P3.3", "hito": "H7",
                "nombre": "Diseño de mapas temáticos",
                "corto": "Diseño de mapas temáticos",
                "detalle": (
                    "Con el estándar definido se componen los mapas temáticos de cada categoría, incorporando leyendas, "
                    "escalas y textos explicativos que permitan a un vecino o a un funcionario entender el mapa sin apoyo técnico."
                ),
                "tecnica": "QGIS · plantillas de composición cartográfica",
                "productos": ["3.2"],
            },
            {
                "id": "P3.4", "hito": "H8",
                "nombre": "Implementación del visor web interactivo",
                "corto": "Implementación del visor web",
                "detalle": (
                    "Se construye el visor web del Atlas con selección de categorías, activación y desactivación de capas, "
                    "consulta de atributos, visualización del ICV, acceso a metadatos ISO 19115, consulta normativa e "
                    "identificación de vacíos territoriales, conectado a servicios espaciales abiertos."
                ),
                "tecnica": "JavaScript · Leaflet · GitHub + Netlify",
                "productos": ["3.3"],
            },
            {
                "id": "P3.5", "hito": "H8",
                "nombre": "Integración de resultados del OE 2 y revisión de la interfaz",
                "corto": "Integración de capas y revisión",
                "detalle": (
                    "Se cargan en el visor la geodatabase y los mapas analíticos del OE 2 y se revisa la interfaz en "
                    "términos de accesibilidad y facilidad de uso antes de abrirla a usuarios externos."
                ),
                "tecnica": "Pruebas internas de navegación · revisión de accesibilidad",
                "productos": ["3.3"],
            },
        ],
        "productos": [
            {"id": "3.1", "nombre": "Estándar cartográfico del Atlas", "corto": "Estándar cartográfico",
             "contenido": "Paleta, simbología, tipografías y plantillas de mapas temáticos.",
             "estado": "programado"},
            {"id": "3.2", "nombre": "Prototipo visual del Atlas", "corto": "Prototipo visual Atlas",
             "contenido": "Identidad gráfica unificada y mapas temáticos finales.",
             "estado": "programado"},
            {"id": "3.3", "nombre": "Versión funcional del visor web", "corto": "Visor web interactivo",
             "contenido": "Atlas interactivo en línea (atlasvgiuah.netlify.app).",
             "estado": "anticipado"},
        ],
    },
    {
        "id": "OE 4",
        "titulo": "Validación participativa y transferibilidad técnica",
        "corto": "Validación participativa\ny transferibilidad",
        "texto": (
            "Validar la usabilidad y pertinencia territorial del Atlas mediante evaluaciones empíricas con actores locales, "
            "generando guías metodológicas y protocolos técnicos de actualización que garanticen la transferibilidad del "
            "proyecto y la autonomía operativa del centro de práctica."
        ),
        "hitos": "Hitos 9 y 10",
        "hitos_corto": "H9–H10",
        "periodo": "Días 45–50 · 05/11 – 12/11/2026",
        "semanas": "Sem 10–11",
        "horas": 36,
        "estado": "programado",
        "procedimientos": [
            {
                "id": "P4.1", "hito": "H9",
                "nombre": "Preparación de la prueba con usuarios",
                "corto": "Preparación prueba usabilidad",
                "detalle": (
                    "Se seleccionan entre 3 y 5 usuarios reales del barrio (vecinos, organizaciones comunitarias y referentes "
                    "institucionales del centro de práctica) y se diseñan tareas tipo y un cuestionario semiestructurado sobre "
                    "legibilidad de los mapas, utilidad de los indicadores, facilidad de navegación e información faltante."
                ),
                "tecnica": "Muestreo intencionado · cuestionario semiestructurado",
                "productos": ["4.1"],
            },
            {
                "id": "P4.2", "hito": "H9",
                "nombre": "Aplicación de las pruebas de usabilidad",
                "corto": "Pruebas de usabilidad en campo",
                "detalle": (
                    "En sesiones guiadas cada participante resuelve las tareas en el Atlas mientras se observa su recorrido y "
                    "se registran dificultades, comentarios y sugerencias."
                ),
                "tecnica": "Prueba de usabilidad observada · registro de campo",
                "productos": ["4.1"],
            },
            {
                "id": "P4.3", "hito": "H9",
                "nombre": "Análisis de observaciones y ajustes iterativos",
                "corto": "Análisis cualitativo y ajustes",
                "detalle": (
                    "Las observaciones se sistematizan y priorizan según su efecto en el uso del Atlas; los cambios se "
                    "implementan de forma iterativa y quedan anotados en un registro de ajustes."
                ),
                "tecnica": "Análisis cualitativo · rediseño iterativo",
                "productos": ["4.1"],
            },
            {
                "id": "P4.4", "hito": "H10",
                "nombre": "Documentación técnica para la reproducibilidad",
                "corto": "Documentación técnica de datos",
                "detalle": (
                    "Se redactan el diccionario de datos, el catálogo de fuentes, la metodología de procesamiento y los "
                    "indicadores, junto con el protocolo para descargar, transformar e integrar nuevas fuentes VGI, de modo "
                    "que el centro de práctica pueda actualizar el Atlas sin depender del practicante."
                ),
                "tecnica": "Redacción técnica · protocolo paso a paso",
                "productos": ["4.3", "4.4"],
            },
            {
                "id": "P4.5", "hito": "H10",
                "nombre": "Manual de uso e informe final",
                "corto": "Manual de uso e informe final",
                "detalle": (
                    "Se elabora el manual de uso dirigido a los usuarios finales del territorio y se redacta el informe de "
                    "práctica profesional que reúne el proceso, los resultados y las recomendaciones."
                ),
                "tecnica": "Redacción técnica y académica",
                "productos": ["4.2", "4.5"],
            },
        ],
        "productos": [
            {"id": "4.1", "nombre": "Versión validada del Atlas", "corto": "Atlas validado con usuarios",
             "contenido": "Atlas ajustado + registro de cambios surgidos de la prueba con usuarios.",
             "estado": "programado"},
            {"id": "4.2", "nombre": "Manual de uso del Atlas", "corto": "Manual de uso del Atlas",
             "contenido": "Guía para usuarios finales del territorio.",
             "estado": "programado"},
            {"id": "4.3", "nombre": "Manual de actualización", "corto": "Manual de actualización",
             "contenido": "Protocolo reproducible de descarga, transformación e integración de fuentes VGI.",
             "estado": "programado"},
            {"id": "4.4", "nombre": "Paquete de datos final", "corto": "Paquete de datos final",
             "contenido": "Geodatabase final, diccionario de datos, catálogo de fuentes, metadatos y metodología.",
             "estado": "programado"},
            {"id": "4.5", "nombre": "Informe de práctica profesional", "corto": "Informe final de práctica",
             "contenido": "Documento final de la práctica.",
             "estado": "programado"},
        ],
    },
]


# ---------------------------------------------------------------------------
# DIAGRAMAS LINEALES (matplotlib)
# ---------------------------------------------------------------------------
plt.rcParams["font.family"] = "DejaVu Sans"


def _caja(ax, x, y, w, h, fc, ec, lw=1.2, radio=0.06):
    ax.add_patch(FancyBboxPatch((x, y), w, h,
                                boxstyle=f"round,pad=0,rounding_size={radio}",
                                fc=fc, ec=ec, lw=lw))


def _flecha(ax, x1, y1, x2, y2, color=HEX_AZUL_SEC, lw=1.6, ms=14):
    ax.add_patch(FancyArrowPatch((x1, y1), (x2, y2), arrowstyle="-|>",
                                 mutation_scale=ms, color=color, lw=lw))


def _leyenda_estados(ax, x, y, claves):
    for k in claves:
        etiqueta, color = ESTADOS[k]
        ax.add_patch(Circle((x, y), 0.065, fc=color, ec="none"))
        ax.text(x + 0.12, y, etiqueta, fontsize=8.2, va="center", color=HEX_NEGRO)
        x += 0.28 + len(etiqueta) * 0.095


def diagrama_global(ruta):
    """OG → OE1 → OE2 → OE3 → OE4 (cadena lineal del plan)."""
    W, H = 10.6, 2.55
    fig, ax = plt.subplots(figsize=(W, H), dpi=220)
    ax.set_xlim(0, W)
    ax.set_ylim(0, H)
    ax.axis("off")

    n = 5
    gap = 0.24
    bw = (W - 0.2 - gap * (n - 1)) / n
    by, bh = 0.62, 1.55

    x = 0.1
    # Objetivo general
    _caja(ax, x, by, bw, bh, HEX_AZUL_UAH, HEX_AZUL_UAH)
    ax.text(x + bw / 2, by + bh - 0.28, "OBJETIVO GENERAL", ha="center", va="center",
            fontsize=8.5, color="white", weight="bold")
    ax.text(x + bw / 2, by + bh / 2 - 0.12, "Prototipo de Atlas VGI\ncentrado en usuarios\npara Barrio Brasil",
            ha="center", va="center", fontsize=8.1, color="white", linespacing=1.25)
    x_prev_end = x + bw

    for oe in OBJETIVOS:
        x = x_prev_end + gap
        _flecha(ax, x_prev_end + 0.02, by + bh / 2, x - 0.02, by + bh / 2)
        _caja(ax, x, by, bw, bh, HEX_AZUL_CLARO, HEX_AZUL_SEC)
        ax.text(x + bw / 2, by + bh - 0.26, oe["id"], ha="center", va="center",
                fontsize=11.0, color=HEX_AZUL_UAH, weight="bold")
        ax.text(x + bw / 2, by + bh / 2 - 0.02, oe["corto"], ha="center", va="center",
                fontsize=7.9, color=HEX_NEGRO, linespacing=1.25)
        ax.text(x + bw / 2, by + 0.2, f'{oe["hitos_corto"]} · {oe["semanas"]} · {oe["horas"]} h',
                ha="center", va="center", fontsize=7.4, color="#475569")
        etiqueta, color = ESTADOS[oe["estado"]]
        ax.add_patch(Circle((x + 0.2, by + bh - 0.26), 0.075, fc=color, ec="none"))
        x_prev_end = x + bw

    ax.text(0.1, 0.36, "Secuencia del plan de trabajo: caracterización → construcción analítica → prototipo → validación y traspaso",
            fontsize=8.3, color="#475569", va="center", style="italic")
    _leyenda_estados(ax, 0.17, 0.08, ["completado", "en_curso", "anticipado", "programado"])
    fig.savefig(ruta, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def diagrama_oe(oe, ruta):
    """Cadena lineal OBJETIVO → PROCEDIMIENTOS (secuencia) → PRODUCTOS."""
    procs = oe["procedimientos"]
    prods = oe["productos"]
    filas = max(len(procs), len(prods))
    paso = 0.56
    W = 10.6
    H = 0.95 + filas * paso + 0.35

    fig, ax = plt.subplots(figsize=(W, H), dpi=220)
    ax.set_xlim(0, W)
    ax.set_ylim(0, H)
    ax.axis("off")

    top = H - 0.45
    base = 0.45
    alto_util = top - base

    # Encabezados de columna
    cols = [(0.10, 1.95, "OBJETIVO"), (2.35, 4.85, "PROCEDIMIENTOS"), (7.55, 2.95, "PRODUCTOS")]
    for cx, cw, ct in cols:
        ax.text(cx + cw / 2, H - 0.2, ct, ha="center", va="center", fontsize=9.5,
                weight="bold", color=HEX_AZUL_UAH)

    # Objetivo
    ox, ow = 0.10, 1.95
    _caja(ax, ox, base, ow, alto_util, HEX_AZUL_UAH, HEX_AZUL_UAH)
    ax.text(ox + ow / 2, base + alto_util / 2 + 0.3, oe["id"], ha="center", va="center",
            fontsize=15, color="white", weight="bold")
    ax.text(ox + ow / 2, base + alto_util / 2 - 0.25, oe["corto"], ha="center", va="center",
            fontsize=8.2, color="white", linespacing=1.25)
    ax.text(ox + ow / 2, base + 0.22, f'{oe["hitos"]} · {oe["horas"]} h', ha="center",
            va="center", fontsize=7.5, color="#CBD5E1")

    # Procedimientos (cadena vertical secuencial)
    px, pw = 2.35, 4.85
    ph = 0.40
    esp_p = alto_util / len(procs)
    centros_p = []
    for i, p in enumerate(procs):
        cy = top - esp_p * (i + 0.5)
        centros_p.append(cy)
        _caja(ax, px, cy - ph / 2, pw, ph, HEX_AZUL_CLARO, HEX_AZUL_SEC, lw=1.0, radio=0.05)
        ax.text(px + 0.15, cy, p["id"], ha="left", va="center", fontsize=8.6,
                weight="bold", color=HEX_AZUL_UAH)
        ax.text(px + 0.65, cy, p["corto"], ha="left", va="center", fontsize=8.3, color=HEX_NEGRO)
        ax.text(px + pw - 0.12, cy, "→ " + ", ".join(p["productos"]), ha="right", va="center",
                fontsize=7.8, color=HEX_AZUL_SEC, weight="bold")
        if i > 0:
            _flecha(ax, px + 0.3, centros_p[i - 1] - ph / 2 - 0.01, px + 0.3, cy + ph / 2 + 0.01,
                    color=HEX_GRIS_LINEA, lw=1.0, ms=8)

    # Productos
    rx, rw = 7.55, 2.95
    rh = 0.40
    esp_r = alto_util / len(prods)
    for i, r in enumerate(prods):
        cy = top - esp_r * (i + 0.5)
        _caja(ax, rx, cy - rh / 2, rw, rh, "white", HEX_AZUL_UAH, lw=1.2, radio=0.05)
        _, color = ESTADOS[r["estado"]]
        ax.add_patch(Circle((rx + 0.17, cy), 0.075, fc=color, ec="none"))
        ax.text(rx + 0.35, cy, r["id"], ha="left", va="center", fontsize=8.6,
                weight="bold", color=HEX_AZUL_UAH)
        nombre = textwrap.shorten(r.get("corto", r["nombre"]), width=28, placeholder="…")
        ax.text(rx + 0.72, cy, nombre, ha="left", va="center", fontsize=8.2, color=HEX_NEGRO)

    # Flechas entre etapas
    ymid = base + alto_util / 2
    _flecha(ax, ox + ow + 0.03, ymid, px - 0.04, ymid, lw=2.0, ms=18)
    _flecha(ax, px + pw + 0.04, ymid, rx - 0.04, ymid, lw=2.0, ms=18)

    # Leyenda de estados
    claves = []
    for r in prods:
        if r["estado"] not in claves:
            claves.append(r["estado"])
    ax.text(0.1, 0.17, "Estado de productos:", fontsize=8.0, va="center", color="#475569", style="italic")
    _leyenda_estados(ax, 1.8, 0.17, claves)

    fig.savefig(ruta, bbox_inches="tight", facecolor="white")
    plt.close(fig)


# ---------------------------------------------------------------------------
# UTILIDADES WORD
# ---------------------------------------------------------------------------
def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>'))


def set_cell_margins(cell, top=70, bottom=70, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(
        f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>'))


def set_table_borders(table, color="CBD5E1", sz="4"):
    table._tbl.tblPr.append(parse_xml(
        f'<w:tblBorders {nsdecls("w")}><w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/><w:left w:val="none"/><w:right w:val="none"/></w:tblBorders>'))


def set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")} w:val="true"/>'))


def keep_row_together(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")} w:val="true"/>'))


def escribir_celda(cell, texto, size=9, bold=False, color=COLOR_NEGRO, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(texto)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    if align is not None:
        p.alignment = align
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def encabezado_tabla(table, headers, widths):
    for ci, h in enumerate(headers):
        c = table.cell(0, ci)
        c.width = widths[ci]
        escribir_celda(c, h, size=9, bold=True, color=COLOR_BLANCO, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_background(c, "003366")
        set_cell_margins(c)
    set_repeat_header(table.rows[0])


def titulo_seccion(doc, texto, nivel=1):
    h = doc.add_heading(level=nivel)
    r = h.add_run(texto)
    r.font.bold = True
    r.font.size = Pt(14 if nivel == 1 else 11.5)
    r.font.color.rgb = COLOR_AZUL_UAH if nivel == 1 else COLOR_AZUL_SEC
    h.paragraph_format.space_before = Pt(14 if nivel == 1 else 10)
    h.paragraph_format.space_after = Pt(6)
    h.paragraph_format.keep_with_next = True
    return h


def parrafo(doc, texto, size=10.5, italic=False, after=6, color=COLOR_NEGRO):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(texto)
    r.font.size = Pt(size)
    r.font.italic = italic
    r.font.color.rgb = color
    return p


def leyenda(doc, texto):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(texto)
    r.font.bold = True
    r.font.size = Pt(9.5)
    r.font.color.rgb = COLOR_AZUL_UAH


def caja_destacada(doc, titulo, texto, fondo="E8F0F8"):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t._tbl.tblPr.append(parse_xml(
        f'<w:tblBorders {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="0" w:color="003366"/>'
        f'<w:top w:val="none"/><w:bottom w:val="none"/><w:right w:val="none"/>'
        f'<w:insideH w:val="none"/><w:insideV w:val="none"/></w:tblBorders>'))
    c = t.cell(0, 0)
    c.width = Inches(6.5)
    set_cell_background(c, fondo)
    set_cell_margins(c, top=110, bottom=110, left=180, right=160)
    c.text = ""
    p1 = c.paragraphs[0]
    p1.paragraph_format.space_after = Pt(3)
    r1 = p1.add_run(titulo)
    r1.font.bold = True
    r1.font.size = Pt(9.5)
    r1.font.color.rgb = COLOR_AZUL_UAH
    p2 = c.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.15
    r2 = p2.add_run(texto)
    r2.font.size = Pt(10.5)
    r2.font.color.rgb = COLOR_NEGRO
    return t


def color_estado(clave):
    hexc = ESTADOS[clave][1].lstrip("#")
    return RGBColor(int(hexc[0:2], 16), int(hexc[2:4], 16), int(hexc[4:6], 16))


# ---------------------------------------------------------------------------
# DOCUMENTO
# ---------------------------------------------------------------------------
def crear_documento():
    # Diagramas
    ruta_global = os.path.join(DIR_DIAGRAMAS, "fig_00_cadena_objetivos.png")
    diagrama_global(ruta_global)
    rutas_oe = []
    for i, oe in enumerate(OBJETIVOS, 1):
        ruta = os.path.join(DIR_DIAGRAMAS, f"fig_0{i}_pipeline_oe{i}.png")
        diagrama_oe(oe, ruta)
        rutas_oe.append(ruta)

    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Inches(1.0)
        s.left_margin = s.right_margin = Inches(1.0)
        # Pie de página
        pf = s.footer.paragraphs[0]
        pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rf = pf.add_run("Atlas VGI · Barrio Brasil – UAH · Informe de Estado de Avance (150 horas)")
        rf.font.size = Pt(8)
        rf.font.color.rgb = COLOR_GRIS

    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    st.font.color.rgb = COLOR_NEGRO

    # ---------------- Encabezado institucional ----------------
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("UNIVERSIDAD ALBERTO HURTADO · FACULTAD DE CIENCIAS SOCIALES\n")
    r.font.bold = True; r.font.size = Pt(10); r.font.color.rgb = COLOR_AZUL_UAH
    r = p.add_run("DEPARTAMENTO DE GEOGRAFÍA · PRÁCTICA PROFESIONAL 2026 S2\n")
    r.font.bold = True; r.font.size = Pt(9.5); r.font.color.rgb = COLOR_AZUL_SEC
    r = p.add_run("Atlas VGI Territorial — Geografía UAH (ANID)")
    r.font.italic = True; r.font.size = Pt(9); r.font.color.rgb = COLOR_GRIS

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("―" * 58)
    r.font.color.rgb = COLOR_AZUL_UAH

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("INFORME DE ESTADO DE AVANCE — 150 HORAS")
    r.font.bold = True; r.font.size = Pt(16); r.font.color.rgb = COLOR_AZUL_UAH
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("Atlas de Información Geográfica Voluntaria para la Gestión Territorial\n"
                  "Caso piloto: Barrio Brasil y entorno del campus UAH (Santiago Centro)")
    r.font.bold = True; r.font.size = Pt(11.5)

    ficha = [
        ("Estudiante practicante", "Juan Pablo Lescano Villegas — Licenciatura en Geografía"),
        ("Coordinador de prácticas", "Alejandro Cortés (UAH)"),
        ("Supervisión", "Prof. Manuel Fuenzalida (UAH) e investigador colaborador Geografía UAH"),
        ("Período de práctica", "01/09/2026 – 12/11/2026 · 300 horas · 50 días hábiles"),
        ("Corte reportado", "150 horas (50 %) · Hitos 1 a 5 concluidos"),
    ]
    t = doc.add_table(rows=len(ficha), cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t, "CBD5E1", "6")
    for i, (k, v) in enumerate(ficha):
        c0, c1 = t.cell(i, 0), t.cell(i, 1)
        c0.width, c1.width = Inches(1.9), Inches(4.6)
        escribir_celda(c0, k, size=9.5, bold=True, color=COLOR_AZUL_UAH)
        escribir_celda(c1, v, size=9.5)
        set_cell_background(c0, "F1F5F9")
        set_cell_margins(c0); set_cell_margins(c1)

    # ---------------- 1. Objetivos y lógica del plan ----------------
    titulo_seccion(doc, "1. Objetivos y lógica del plan de trabajo")
    caja_destacada(doc, "OBJETIVO GENERAL", OBJETIVO_GENERAL)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    parrafo(doc,
            "El plan se organiza en cuatro objetivos específicos que se encadenan: primero se caracteriza el territorio "
            "y se diseña el Atlas (OE 1), luego se construye y evalúa la base de datos (OE 2), después se diseña y publica "
            "el prototipo (OE 3) y, finalmente, se valida con usuarios y se deja documentado para su continuidad (OE 4). "
            "Cada objetivo se desarrolla en este informe como una cadena Objetivo → Procedimientos → Productos: los "
            "procedimientos describen qué se hizo o se hará y con qué técnica, y cada uno indica el producto que alimenta.")

    doc.add_picture(ruta_global, width=Inches(6.5))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    leyenda(doc, "Figura 1. Cadena de objetivos del plan de trabajo (hitos, semanas, horas y estado al corte).")

    leyenda(doc, "Tabla 1. Distribución de objetivos, hitos, horas y estado de avance")
    t = doc.add_table(rows=len(OBJETIVOS) + 2, cols=5)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t)
    widths = [Inches(0.6), Inches(2.6), Inches(1.6), Inches(0.6), Inches(1.1)]
    encabezado_tabla(t, ["OE", "Objetivo específico", "Hitos y período", "Horas", "Estado"], widths)
    for i, oe in enumerate(OBJETIVOS, 1):
        fila = t.rows[i]
        bg = "FFFFFF" if i % 2 else "F8FAFC"
        vals = [oe["id"], oe["titulo"], f'{oe["hitos"]}\n{oe["periodo"]}', str(oe["horas"]), ESTADOS[oe["estado"]][0]]
        for ci, v in enumerate(vals):
            c = fila.cells[ci]
            c.width = widths[ci]
            escribir_celda(c, v, size=8.8,
                           bold=(ci in (0, 4)),
                           color=(COLOR_AZUL_UAH if ci == 0 else (color_estado(oe["estado"]) if ci == 4 else COLOR_NEGRO)),
                           align=(WD_ALIGN_PARAGRAPH.CENTER if ci in (0, 3, 4) else None))
            set_cell_background(c, bg)
            set_cell_margins(c)
    tot = t.rows[-1]
    for ci, v in enumerate(["", "Total planificado", "10 hitos · 50 días hábiles", "300", "150 h ejecutadas"]):
        c = tot.cells[ci]
        c.width = widths[ci]
        escribir_celda(c, v, size=8.8, bold=True, color=COLOR_AZUL_UAH,
                       align=(WD_ALIGN_PARAGRAPH.CENTER if ci in (0, 3, 4) else None))
        set_cell_background(c, "E8F0F8")
        set_cell_margins(c)

    # ---------------- 2–5. Un apartado por objetivo ----------------
    for n, (oe, ruta) in enumerate(zip(OBJETIVOS, rutas_oe), start=2):
        doc.add_page_break()
        titulo_seccion(doc, f'{n}. {oe["id"]} — {oe["titulo"]}')
        caja_destacada(doc, f'OBJETIVO ESPECÍFICO {oe["id"][-1]}  ·  {oe["hitos"]}  ·  {oe["periodo"]}  ·  '
                            f'{oe["horas"]} h  ·  {ESTADOS[oe["estado"]][0].upper()}', oe["texto"])
        doc.add_paragraph().paragraph_format.space_after = Pt(0)

        doc.add_picture(ruta, width=Inches(6.5))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        leyenda(doc, f'Figura {n}. Cadena Objetivo → Procedimientos → Productos del {oe["id"]}.')

        # Plan metodológico
        titulo_seccion(doc, f'{n}.1 Plan metodológico (procedimientos)', nivel=2)
        for pr in oe["procedimientos"]:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(f'{pr["id"]}  {pr["nombre"]}')
            r.font.bold = True; r.font.size = Pt(10.5); r.font.color.rgb = COLOR_AZUL_SEC
            r = p.add_run(f'   ({pr["hito"]})')
            r.font.size = Pt(9); r.font.color.rgb = COLOR_GRIS

            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run(pr["detalle"])
            r.font.size = Pt(10)

            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_after = Pt(7)
            r = p.add_run("Técnica: "); r.font.bold = True; r.font.size = Pt(9); r.font.color.rgb = COLOR_GRIS
            r = p.add_run(pr["tecnica"]); r.font.size = Pt(9); r.font.color.rgb = COLOR_GRIS
            r = p.add_run("     Producto que alimenta: "); r.font.bold = True; r.font.size = Pt(9); r.font.color.rgb = COLOR_AZUL_UAH
            r = p.add_run(", ".join(pr["productos"])); r.font.bold = True; r.font.size = Pt(9); r.font.color.rgb = COLOR_AZUL_UAH

        # Productos
        titulo_seccion(doc, f'{n}.2 Productos específicos', nivel=2)
        leyenda(doc, f'Tabla {n}. Productos del {oe["id"]} y procedimientos que los generan')
        t = doc.add_table(rows=len(oe["productos"]) + 1, cols=5)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(t)
        widths = [Inches(0.45), Inches(1.65), Inches(2.5), Inches(0.85), Inches(1.05)]
        encabezado_tabla(t, ["N°", "Producto", "Contenido", "Generado por", "Estado"], widths)
        for i, prod in enumerate(oe["productos"], 1):
            fila = t.rows[i]
            keep_row_together(fila)
            bg = "FFFFFF" if i % 2 else "F8FAFC"
            origen = ", ".join(pr["id"] for pr in oe["procedimientos"] if prod["id"] in pr["productos"])
            vals = [prod["id"], prod["nombre"], prod["contenido"], origen, ESTADOS[prod["estado"]][0]]
            for ci, v in enumerate(vals):
                c = fila.cells[ci]
                c.width = widths[ci]
                escribir_celda(c, v, size=8.8,
                               bold=(ci in (0, 1, 4)),
                               color=(COLOR_AZUL_UAH if ci == 0 else (color_estado(prod["estado"]) if ci == 4 else COLOR_NEGRO)),
                               align=(WD_ALIGN_PARAGRAPH.CENTER if ci in (0, 3, 4) else None))
                set_cell_background(c, bg)
                set_cell_margins(c)

    # ---------------- 6. Próximos pasos ----------------
    doc.add_page_break()
    titulo_seccion(doc, "6. Próximos pasos (150 horas restantes)")
    parrafo(doc, "Las actividades pendientes corresponden exclusivamente a los hitos ya comprometidos en el plan de trabajo:")
    pasos = [
        ("Hito 6 · OE 2 (Sem 6–7):", "cerrar el análisis territorial (P2.6) y entregar los mapas analíticos (producto 2.2)."),
        ("Hito 7 · OE 3 (Sem 8):", "formalizar el estándar cartográfico y los mapas temáticos (P3.1–P3.3; productos 3.1 y 3.2)."),
        ("Hito 8 · OE 3 (Sem 8–10):", "integrar los mapas analíticos en el visor y revisar su usabilidad (P3.4–P3.5; producto 3.3)."),
        ("Hito 9 · OE 4 (Sem 10–11):", "aplicar la prueba con 3–5 usuarios del barrio y realizar los ajustes (P4.1–P4.3; producto 4.1)."),
        ("Hito 10 · OE 4 (Sem 11):", "entregar manuales, paquete de datos final e informe de práctica (P4.4–P4.5; productos 4.2–4.5)."),
    ]
    for k, v in pasos:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(k + " "); r.font.bold = True; r.font.color.rgb = COLOR_AZUL_SEC
        p.add_run(v)

    # Firma
    doc.add_paragraph().paragraph_format.space_after = Pt(18)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Juan Pablo Lescano Villegas\n")
    r.font.bold = True; r.font.size = Pt(11); r.font.color.rgb = COLOR_AZUL_UAH
    r = p.add_run("Estudiante de Geografía · Practicante Departamento de Geografía UAH\nUniversidad Alberto Hurtado")
    r.font.size = Pt(9.5); r.font.color.rgb = COLOR_GRIS

    doc.save(OUTPUT_DOCX)
    print(f"✅ Documento Word generado: {OUTPUT_DOCX}")
    print(f"   Diagramas en: {DIR_DIAGRAMAS}")


if __name__ == "__main__":
    crear_documento()
