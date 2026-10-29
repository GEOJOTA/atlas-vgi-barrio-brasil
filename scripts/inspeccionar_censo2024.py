#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
inspeccionar_censo2024.py
Inspecciona los archivos GPKG y Excel del Censo 2024 (Región Metropolitana).
"""

import sqlite3
import os
import sys
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

base_dir = r"C:\Users\PT1261256\Desktop\16_09_2026\Base de datos VGI\CENSO 2024"
gpkg_files = [f for f in os.listdir(base_dir) if f.endswith(".gpkg")]

print("=" * 80)
print("INSPECCIÓN DE CAPAS CENSO 2024 — REGIÓN METROPOLITANA")
print("=" * 80)

for g in sorted(gpkg_files):
    path = os.path.join(base_dir, g)
    size_mb = os.path.getsize(path) / (1024 * 1024)
    conn = sqlite3.connect(path)
    cur = conn.cursor()
    
    cur.execute("SELECT table_name, data_type, identifier, srs_id, min_x, min_y, max_x, max_y FROM gpkg_contents")
    contents = cur.fetchall()
    
    print(f"\n📁 Archivo: {g} ({size_mb:.2f} MB)")
    for table_name, data_type, identifier, srs_id, min_x, min_y, max_x, max_y in contents:
        cur.execute(f'SELECT count(*) FROM "{table_name}"')
        count = cur.fetchone()[0]
        cur.execute(f'PRAGMA table_info("{table_name}")')
        cols = [c[1] for c in cur.fetchall()]
        print(f"   Tabla: {table_name}")
        print(f"   Tipo: {data_type} | SRID: {srs_id}")
        print(f"   BBox: [{min_x}, {min_y}, {max_x}, {max_y}]")
        print(f"   Total Entidades: {count:,}")
        print(f"   Columnas ({len(cols)}): {', '.join(cols[:20])}")
        
        # Muestra de 1 registro
        cur.execute(f'SELECT * FROM "{table_name}" LIMIT 1')
        sample = cur.fetchone()
        if sample:
            sample_dict = {cols[i]: sample[i] for i in range(min(len(cols), 10)) if cols[i] != 'geom'}
            print(f"   Muestra de datos: {sample_dict}")
            
    conn.close()

# Inspeccionar Excel
xlsx_path = os.path.join(base_dir, "diccionario_variables_glosas_censo2024.xlsx")
if os.path.exists(xlsx_path):
    print("\n" + "=" * 80)
    print("DICCIONARIO DE VARIABLES GLOSAS CENSO 2024")
    print("=" * 80)
    wb = openpyxl.load_workbook(xlsx_path, read_only=True)
    print(f"Hojas en Excel ({len(wb.sheetnames)}): {wb.sheetnames}")
    for sheet in wb.sheetnames[:5]:
        ws = wb[sheet]
        rows = list(ws.iter_rows(max_row=5, values_only=True))
        print(f"\nHoja: {sheet}")
        for r in rows:
            print("  ", [str(x)[:30] if x is not None else "" for x in r[:8]])
    wb.close()
